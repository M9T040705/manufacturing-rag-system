"""
文档加载器
支持 PDF、Word、PPT、TXT、图片（OCR）等格式
统一输出 LangChain Document 对象
"""
import os
from pathlib import Path
from typing import List, Optional

from langchain_core.documents import Document
from loguru import logger

from config.settings import settings


class DocumentLoader:
    """多格式文档统一加载器"""

    SUPPORTED_EXTENSIONS = {
        ".pdf": "_load_pdf",
        ".docx": "_load_docx",
        ".doc": "_load_doc",
        ".pptx": "_load_pptx",
        ".txt": "_load_txt",
        ".png": "_load_image",
        ".jpg": "_load_image",
        ".jpeg": "_load_image",
    }

    def __init__(self):
        self.ocr_processor = None
        if settings.ocr_enabled:
            try:
                from .ocr import OCRProcessor
                self.ocr_processor = OCRProcessor()
                logger.info("OCR 处理器初始化成功")
            except Exception as e:
                logger.warning(f"OCR 初始化失败，图片将无法识别: {e}")

    def load(self, file_path: str) -> List[Document]:
        """
        加载单个文档文件

        Args:
            file_path: 文件绝对路径

        Returns:
            Document 列表，每个元素代表一页/一段
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支持的文件格式: {ext}")

        loader_method = getattr(self, self.SUPPORTED_EXTENSIONS[ext])
        logger.info(f"开始加载文档: {path.name} ({ext})")

        try:
            docs = loader_method(file_path)
            # 统一添加元数据
            for doc in docs:
                doc.metadata["source"] = str(path)
                doc.metadata["filename"] = path.name
                doc.metadata["file_type"] = ext
            logger.info(f"文档加载完成: {path.name}, 共 {len(docs)} 页/段")
            return docs
        except Exception as e:
            logger.error(f"文档加载失败 {path.name}: {e}")
            raise

    def load_directory(self, dir_path: str, recursive: bool = True) -> List[Document]:
        """批量加载目录下所有支持的文档"""
        all_docs = []
        path = Path(dir_path)

        pattern = "**/*" if recursive else "*"
        for file_path in path.glob(pattern):
            if file_path.is_file() and file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                try:
                    docs = self.load(str(file_path))
                    all_docs.extend(docs)
                except Exception as e:
                    logger.warning(f"跳过文件 {file_path.name}: {e}")

        logger.info(f"目录加载完成: {dir_path}, 共 {len(all_docs)} 个文档片段")
        return all_docs

    def _load_pdf(self, file_path: str) -> List[Document]:
        """加载 PDF，优先文本提取，扫描件走 OCR"""
        import fitz  # PyMuPDF

        docs = []
        doc = fitz.open(file_path)

        for page_num, page in enumerate(doc):
            text = page.get_text().strip()

            # 如果文本过少，可能是扫描件，走 OCR
            if len(text) < 50 and self.ocr_processor:
                logger.debug(f"PDF 第 {page_num + 1} 页文本过少，尝试 OCR")
                pix = page.get_pixmap(dpi=300)
                img_bytes = pix.tobytes("png")
                text = self.ocr_processor.recognize_bytes(img_bytes)

            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={"page": page_num + 1, "total_pages": len(doc)}
                ))

        doc.close()
        return docs

    def _load_docx(self, file_path: str) -> List[Document]:
        """加载 Word 文档，按标题层级分段"""
        from docx import Document as DocxDocument

        docx_doc = DocxDocument(file_path)
        docs = []
        current_heading = ""
        current_text = []

        for para in docx_doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()

            if not text:
                continue

            if style.startswith("Heading"):
                # 保存上一段
                if current_text:
                    docs.append(Document(
                        page_content="\n".join(current_text),
                        metadata={"heading": current_heading}
                    ))
                current_heading = text
                current_text = []
            else:
                current_text.append(text)

        # 最后一段
        if current_text:
            docs.append(Document(
                page_content="\n".join(current_text),
                metadata={"heading": current_heading}
            ))

        return docs if docs else [Document(page_content="\n".join(
            p.text for p in docx_doc.paragraphs if p.text.strip()
        ))]

    def _load_doc(self, file_path: str) -> List[Document]:
        """加载旧版 .doc（通过 antiword 或 textract）"""
        # 简化处理：尝试用 textract，否则提示转换
        try:
            import textract
            text = textract.process(file_path).decode("utf-8", errors="ignore")
            return [Document(page_content=text)]
        except ImportError:
            logger.warning("textract 未安装，.doc 文件将尝试以文本读取")
            return self._load_txt(file_path)

    def _load_pptx(self, file_path: str) -> List[Document]:
        """加载 PPT，按幻灯片分页"""
        from pptx import Presentation

        prs = Presentation(file_path)
        docs = []

        for slide_num, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())

            if texts:
                docs.append(Document(
                    page_content="\n".join(texts),
                    metadata={"slide": slide_num + 1, "total_slides": len(prs.slides)}
                ))

        return docs

    def _load_txt(self, file_path: str) -> List[Document]:
        """加载纯文本，自动检测编码"""
        import chardet

        with open(file_path, "rb") as f:
            raw = f.read()

        detected = chardet.detect(raw)
        encoding = detected.get("encoding", "utf-8") or "utf-8"

        try:
            text = raw.decode(encoding, errors="ignore")
        except Exception:
            text = raw.decode("utf-8", errors="ignore")

        return [Document(page_content=text)]

    def _load_image(self, file_path: str) -> List[Document]:
        """加载图片，通过 OCR 提取文字"""
        if not self.ocr_processor:
            raise RuntimeError("OCR 未启用，无法处理图片")

        text = self.ocr_processor.recognize_file(file_path)
        if not text.strip():
            logger.warning(f"图片未识别到文字: {file_path}")
            return []

        return [Document(page_content=text, metadata={"ocr": True})]
